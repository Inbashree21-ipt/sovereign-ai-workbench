"""
Spreadsheet and Tabular Data Analysis Tool

Supports:
- CSV files
- TXT files
- Excel XLSX/XLSM files
- Column sums
- Mean
- Minimum
- Maximum
- Optional validation against expected totals

Expected totals are OPTIONAL.
If the user only asks for totals, the tool calculates the
actual totals without inventing expected values.
"""

import csv
from pathlib import Path
from typing import Dict, Any, List

import openpyxl

from tools.base import BaseTool, ToolResult
from tools.file_ops import _resolve_safe_path


class AnalyzeSpreadsheetTool(BaseTool):

    name = "analyze_spreadsheet"

    description = (
        "Analyzes a CSV or Excel spreadsheet in the workspace. "
        "Calculates sums, averages, minimums and maximums for numeric columns. "
        "Expected totals are optional and must only be provided when explicitly known."
    )

    parameters = {
        "type": "object",
        "properties": {

            "file_path": {
                "type": "string",
                "description": (
                    "Relative path to the CSV, TXT, XLSX or XLSM file "
                    "inside the workspace."
                ),
            },

            "columns": {
                "type": "array",
                "items": {
                    "type": "string"
                },
                "description": (
                    "Optional list of numeric column names to analyze. "
                    "If omitted, all numeric columns are analyzed."
                ),
            },

            "expected_totals": {
                "type": "object",
                "description": (
                    "OPTIONAL. Only provide this when the user explicitly "
                    "supplies expected totals. Never invent expected values."
                ),
            },
        },

        "required": ["file_path"],
    }

    # ---------------------------------------------------------
    # READ TABLE
    # ---------------------------------------------------------

    def _read_table(self, target_path: Path) -> List[Dict[str, Any]]:
        """
        Reads CSV or Excel data and converts it into
        a list of dictionaries.
        """

        rows = []

        suffix = target_path.suffix.lower()

        # -----------------------------------------------------
        # Excel
        # -----------------------------------------------------

        if suffix in [".xlsx", ".xlsm"]:

            wb = openpyxl.load_workbook(
                str(target_path),
                data_only=True
            )

            ws = wb.active

            iter_rows = list(
                ws.iter_rows(values_only=True)
            )

            if not iter_rows:
                return []

            headers = []

            for i, header in enumerate(iter_rows[0]):

                if header is None:
                    headers.append(f"Col_{i}")

                else:
                    headers.append(
                        str(header).strip()
                    )

            for row in iter_rows[1:]:

                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                row_dict = {}

                for i, header in enumerate(headers):

                    value = (
                        row[i]
                        if i < len(row)
                        else None
                    )

                    row_dict[header] = value

                rows.append(row_dict)

        # -----------------------------------------------------
        # CSV / TXT
        # -----------------------------------------------------

        elif suffix in [".csv", ".txt"]:

            with open(
                target_path,
                "r",
                encoding="utf-8-sig",
                errors="replace"
            ) as file:

                reader = csv.DictReader(file)

                for row in reader:

                    row_dict = {}

                    for key, value in row.items():

                        cleaned_key = (
                            key.strip()
                            if key
                            else "Col"
                        )

                        cleaned_value = (
                            value.strip()
                            if isinstance(value, str)
                            else value
                        )

                        row_dict[cleaned_key] = (
                            cleaned_value
                        )

                    rows.append(row_dict)

        else:

            raise ValueError(
                f"Unsupported spreadsheet format '{suffix}'. "
                "Use .csv, .txt, .xlsx or .xlsm."
            )

        return rows

    # ---------------------------------------------------------
    # CONVERT VALUE TO NUMBER
    # ---------------------------------------------------------

    def _to_number(self, value: Any):

        if value is None:
            return None

        if isinstance(value, (int, float)):

            return float(value)

        if isinstance(value, str):

            value = (
                value
                .replace("$", "")
                .replace("₹", "")
                .replace(",", "")
                .strip()
            )

            if not value:
                return None

            try:
                return float(value)

            except ValueError:
                return None

        return None

    # ---------------------------------------------------------
    # RUN TOOL
    # ---------------------------------------------------------

    def run(self, **kwargs) -> ToolResult:

        file_path_str = kwargs.get(
            "file_path",
            ""
        )

        requested_columns = kwargs.get(
            "columns",
            None
        )

        expected_totals = kwargs.get(
            "expected_totals",
            None
        )

        try:

            # -------------------------------------------------
            # Resolve safe workspace path
            # -------------------------------------------------

            target_path = _resolve_safe_path(
                file_path_str
            )

            if not target_path.exists():

                return ToolResult(
                    success=False,
                    output="",
                    error=(
                        f"Spreadsheet '{file_path_str}' "
                        "not found in workspace."
                    ),
                )

            # -------------------------------------------------
            # Read data
            # -------------------------------------------------

            rows = self._read_table(
                target_path
            )

            if not rows:

                return ToolResult(
                    success=True,
                    output=(
                        f"Spreadsheet '{file_path_str}' "
                        "is empty."
                    ),
                    data={
                        "rows": 0,
                        "columns": 0,
                    },
                )

            # -------------------------------------------------
            # Determine columns
            # -------------------------------------------------

            all_columns = list(
                rows[0].keys()
            )

            row_count = len(rows)

            if requested_columns:

                candidate_columns = (
                    requested_columns
                )

            else:

                candidate_columns = (
                    all_columns
                )

            # -------------------------------------------------
            # Calculate statistics
            # -------------------------------------------------

            stats = {}

            for column in candidate_columns:

                if column not in all_columns:
                    continue

                numeric_values = []

                for row in rows:

                    value = self._to_number(
                        row.get(column)
                    )

                    if value is not None:

                        numeric_values.append(
                            value
                        )

                if not numeric_values:
                    continue

                total = sum(
                    numeric_values
                )

                count = len(
                    numeric_values
                )

                mean = total / count

                minimum = min(
                    numeric_values
                )

                maximum = max(
                    numeric_values
                )

                stats[column] = {

                    "count": count,

                    "sum": round(
                        total,
                        2
                    ),

                    "mean": round(
                        mean,
                        2
                    ),

                    "min": round(
                        minimum,
                        2
                    ),

                    "max": round(
                        maximum,
                        2
                    ),
                }

            # -------------------------------------------------
            # Optional expected-total validation
            # -------------------------------------------------

            validations = []

            has_discrepancy = False

            if (
                expected_totals
                and isinstance(
                    expected_totals,
                    dict
                )
            ):

                for (
                    column_name,
                    expected_value
                ) in expected_totals.items():

                    if column_name in stats:

                        actual_value = (
                            stats[column_name]["sum"]
                        )

                        difference = round(
                            actual_value
                            - float(expected_value),
                            2
                        )

                        matches = (
                            abs(difference)
                            < 0.01
                        )

                        if not matches:
                            has_discrepancy = True

                        validations.append({

                            "column": column_name,

                            "expected": float(
                                expected_value
                            ),

                            "actual": actual_value,

                            "discrepancy": difference,

                            "is_valid": matches,
                        })

                    else:

                        validations.append({

                            "column": column_name,

                            "expected": float(
                                expected_value
                            ),

                            "error": (
                                f"Numeric column "
                                f"'{column_name}' "
                                "not found."
                            ),

                            "is_valid": False,
                        })

            # -------------------------------------------------
            # Build output
            # -------------------------------------------------

            output_lines = [

                f"Spreadsheet Analysis for "
                f"'{file_path_str}':",

                (
                    f"- Total Rows: {row_count}, "
                    f"Total Columns: "
                    f"{len(all_columns)}"
                ),

                (
                    "- Columns Available: "
                    + ", ".join(all_columns)
                ),

                "",

                "Summary Calculations:",
            ]

            # -------------------------------------------------
            # Add statistics
            # -------------------------------------------------

            for column, data in stats.items():

                output_lines.append(

                    f"  * {column} -> "
                    f"Sum: {data['sum']:,.2f} | "
                    f"Mean: {data['mean']:,.2f} | "
                    f"Min: {data['min']:,.2f} | "
                    f"Max: {data['max']:,.2f} "
                    f"(Count: {data['count']})"
                )

            # -------------------------------------------------
            # Add validation ONLY when expected totals
            # exist
            # -------------------------------------------------

            if validations:

                output_lines.append(
                    ""
                )

                output_lines.append(
                    "Verification Results "
                    "against Expected Totals:"
                )

                for validation in validations:

                    if validation.get(
                        "is_valid"
                    ):

                        status = (
                            "MATCH (VALID)"
                        )

                    else:

                        status = (
                            "MISMATCH "
                            "(FLAGGED DISCREPANCY)"
                        )

                    if (
                        "discrepancy"
                        in validation
                    ):

                        output_lines.append(

                            f"  * [{status}] "
                            f"Column "
                            f"'{validation['column']}': "
                            f"Expected="
                            f"{validation['expected']:,.2f}, "
                            f"Actual="
                            f"{validation['actual']:,.2f} "
                            f"(Variance: "
                            f"{validation['discrepancy']:+,.2f})"
                        )

                    else:

                        output_lines.append(

                            f"  * [{status}] "
                            f"{validation.get('error')}"
                        )

            # -------------------------------------------------
            # Return result
            # -------------------------------------------------

            return ToolResult(

                success=True,

                output="\n".join(
                    output_lines
                ),

                data={

                    "rows": row_count,

                    "columns": len(
                        all_columns
                    ),

                    "stats": stats,

                    "validations": validations,

                    "has_discrepancy": (
                        has_discrepancy
                    ),
                },
            )

        except Exception as e:

            return ToolResult(

                success=False,

                output="",

                error=(
                    "Spreadsheet analysis error: "
                    f"{str(e)}"
                ),
            )

